from pathlib import Path
from shutil import copyfile
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT=Path(__file__).resolve().parent
TEMPLATE=Path('/Users/shkodranasani/.codex/plugins/cache/openai-curated-remote/openai-templates/0.1.1/skills/artifact-template-legal-memorandum/assets/reference.docx')
OUT=ROOT/'Report_Sem01_1.1_Matrice_regole_norme_CH.docx'; MD=ROOT/'docs/REPORT_1.1_RULES_CH.md'
R=[
('R01','Priorità del segnale luminoso','Vorrang des Lichtsignals','Il segnale luminoso domina la scelta del movimento.','SSV 68.1','LEGGE','corrispondenza diretta'),
('R02','Rosso arresta','Rot bedeutet Halt','Con rosso nessun veicolo entra.','SSV 68.1bis','LEGGE','diretta'),
('R03','Verde autorizza','Grün gibt den Verkehr frei','Entra soltanto il gruppo autorizzato.','SSV 68.2','LEGGE','restano gli obblighi di precedenza'),
('R04','Giallo dopo verde','Gelb nach Grün','Verde passa sempre da giallo prima del rosso.','SSV 68.4 lett. a','LEGGE','diretta'),
('R05','Rosso più giallo prima del verde','Rot und Gelb vor Grün','La fase prepara il verde con rosso e giallo.','SSV 68.4 lett. b','LEGGE','diretta'),
('R06','Ordine delle luci','Anordnung der Lichter','Rosso alto, giallo centro, verde basso.','SSV 70.5','LEGGE','diretta'),
('R07','Frecce direzionali','Richtungspfeile','Segnale e autorizzazione valgono per la direzione indicata.','SSV 68.3, 68.5, 68.9','LEGGE','diretta'),
('R08','Esclusione dei conflitti','Ausschluss von Konflikten','La matrice vieta verdi incompatibili.','SSV 71.3','LEGGE','la geometria va validata sul sito'),
('R09','Frecce verdi protette','Geschützte Grünpfeile','Senza giallo lampeggiante escludono pedoni e traffico conflittuale.','SSV 71.3','LEGGE','diretta'),
('R10','Posizione dei semafori','Standort der Lichtsignale','Ogni gruppo è associato chiaramente alla corsia.','SSV 71.1–1bis','LEGGE','quote reali da progettare'),
('R11','Corsie e preselezione','Fahrstreifen und Einspurpfeile','Corsia sinistra e corsia diritto/destra; deviazione registrata.','SSV 74.1–2','LEGGE','la geometria è specifica del sito'),
('R12','Linea di arresto','Haltelinie','I veicoli attendono prima della linea bianca.','SSV 75.1','LEGGE','diretta'),
('R13','Verde pedonale','Fussgängergrün','Il pedone entra solo con simbolo verde.','SSV 68.7; 70.7','LEGGE','diretta'),
('R14','Sgombero pedonale','Fussgängerräumung','A fine verde nessun nuovo ingresso; chi attraversa esce subito.','SSV 68.7','LEGGE','durata da calcolo tecnico'),
('R15','Chiamata pedonale memorizzata','Gespeicherte Anforderung','La richiesta resta fino a una fase sicura.','Norme VSS e progetto autorizzato','TECNICA','nessun singolo articolo prescrive l’algoritmo'),
('R16','Traffico compatibile col pedone','Verträglicher Verkehr','Avanzano soltanto movimenti che non incontrano le strisce attive.','SSV 71.3','LEGGE','serve esclusione geometrica verificata'),
('R17','Biciclette','Fahrräder','Velocità propria; eventuali luci dedicate usano il simbolo bici.','SSV 68.8; 70.7','LEGGE','il tracciamento lento è scelta tecnica'),
('R18','Zona protetta','Schutzzone','Nessuna fase conflittuale parte finché la zona è occupata.','SSV 71.3; norme VSS sugli intertempi','TECNICA','principio legale più realizzazione tecnica'),
('R19','Sgombero adattivo','Adaptive Räumzeit','Velocità, lunghezza e coda mantengono occupata la zona.','Norme VSS; ASTRA 11001 per impianti provvisori','TECNICA','valori e margini specifici del sito'),
('R20','Bus camion e articolati','Busse Lastwagen Sattelschlepper','L’uscita considera la parte posteriore e la lunghezza.','Analisi dei rischi e norme tecniche','PROGETTO','barriera prudenziale, non algoritmo imposto dalla SSV'),
('R21','Domanda e rotazione','Anforderung und Rotation','Scelta per domanda e compatibilità, con rotazione in parità.','Norme VSS e programma autorizzato','TECNICA','da dimensionare sui flussi reali'),
('R22','Guasto critico','Kritischer Fehler','Stato sicuro memorizzato fino a ripristino esplicito.','Analisi di sicurezza e norme applicabili','PROGETTO','fail-safe non certificato dal solo software'),
('R23','Perdita sensore','Sensorausfall','Mancata conferma non equivale a zona libera.','Analisi di sicurezza e norme tecniche','PROGETTO','regola prudenziale fail-safe'),
('R24','Direzione imprevista','Unerwartete Richtung','Anomalia registrata e protezione del movimento conflittuale.','SSV 74.2; SVG prudenza generale','LEGGE','il recupero software è barriera aggiuntiva'),
('R25','Cantiere autorizzato','Bewilligte Baustelle','Modalità cantiere separata dall’incrocio.','SSV 9, 80, 81, 107','LEGGE','posa e ordinanze richiedono autorità/ASTRA'),
('R26','Cantiere madre figlio','Master Slave Baustelle','Un senso verde; inversione solo dopo corridoio libero.','SSV 71.3; ASTRA 11001; VSS','TECNICA','architettura di progetto'),
('R27','Perdita radio','Funkausfall','Blocca nuove autorizzazioni e genera guasto.','Analisi di sicurezza e norme tecniche','PROGETTO','da validare su hardware'),
('R28','Giallo lampeggiante','Gelbes Blinklicht','Usato solo nei casi ammessi, non come ciclo normale.','SSV 68.6; 70.1','LEGGE','diretta'),
('R29','Segnalazione del cantiere','Baustellensignalisation','La versione reale richiede segnale 1.14 e dispositivi temporanei.','SSV 9, 80, 82','LEGGE','non ancora completa nella simulazione'),
('R30','Camera e radar','Kamera und Radar','Stima presenza e velocità; non serve conservare immagini identificative.','LPD/DSG; indicazioni EDÖB','LEGGE','suolo pubblico: base legale e autorità; minimizzazione dati'),
('R31','Parametri del sito','Standortparameter','Tempi, distanze e timeout configurabili.','Norme VSS, progetto e autorizzazione','TECNICA','i valori demo non valgono automaticamente su strada'),
('R32','Registro e lingue','Protokoll und Sprachen','Interfaccia IT/DE e registro rendono verificabili le decisioni.','Documentazione e collaudo','PROGETTO','qualità e tracciabilità, non articolo stradale'),
]
S=[('SSV ufficiale italiano','https://www.fedlex.admin.ch/eli/cc/1979/1961_1961_1961/it'),('SSV amtlich Deutsch','https://www.fedlex.admin.ch/eli/cc/1979/1961_1961_1961/de'),('SVG RS 741.01','https://www.fedlex.admin.ch/eli/cc/1959/679_705_685/it'),('VRV ONC RS 741.11','https://www.fedlex.admin.ch/eli/cc/1962/1364_1409_1420/it'),('ASTRA segnali','https://www.astra.admin.ch/it/segnali'),('EDÖB videosorveglianza spazio pubblico','https://www.edoeb.admin.ch/it/videosorveglianza-dello-spazio-pubblico-da-parte-di-privati')]

def shade(c,fill):
 e=OxmlElement('w:shd'); e.set(qn('w:fill'),fill); c._tc.get_or_add_tcPr().append(e)
def borders(t):
 b=OxmlElement('w:tblBorders')
 for n in ('top','left','bottom','right','insideH','insideV'):
  e=OxmlElement('w:'+n); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); b.append(e)
 t._tbl.tblPr.append(b)
def margin(c):
 m=OxmlElement('w:tcMar')
 for n in ('top','start','bottom','end'):
  e=OxmlElement('w:'+n); e.set(qn('w:w'),'100'); e.set(qn('w:type'),'dxa'); m.append(e)
 c._tc.get_or_add_tcPr().append(m)
def link(p,label,url):
 rid=p.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True); h=OxmlElement('w:hyperlink'); h.set(qn('r:id'),rid); r=OxmlElement('w:r'); rp=OxmlElement('w:rPr'); c=OxmlElement('w:color'); c.set(qn('w:val'),'0563C1'); rp.append(c); u=OxmlElement('w:u'); u.set(qn('w:val'),'single'); rp.append(u); r.append(rp); x=OxmlElement('w:t'); x.text=label; r.append(x); h.append(r); p._p.append(h)

copyfile(TEMPLATE,OUT); d=Document(OUT); body=d._element.body; sp=body.sectPr
for section in d.sections:
 for index,hp in enumerate(section.header.paragraphs):
  hp.text='SEM.01  REPORT 1.1' if index==0 else ''; hp.alignment=WD_ALIGN_PARAGRAPH.CENTER
  for run in hp.runs: run.font.name='Arial'; run.font.size=Pt(8); run.font.color.rgb=RGBColor(130,130,130)
for x in list(body):
 if x is not sp: body.remove(x)
d.styles['Normal'].font.name='Garamond'; d.styles['Normal'].font.size=Pt(10.5)
p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; z=p.add_run('REVISIONE TECNICA PUBBLICA'); z.font.name='Arial'; z.font.size=Pt(9); z.font.color.rgb=RGBColor(120,120,120)
p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; z=p.add_run('MEMORANDUM SEM 01'); z.bold=True; z.font.size=Pt(26)
for a,b in [('Destinatario','Team Sem.01 e progettista dell’impianto'),('Autore','Progetto Sem.01'),('Data','5 settembre 2026'),('Oggetto','Regole software e quadro svizzero')]: p=d.add_paragraph(); p.add_run(a+':\t').bold=True; p.add_run(b)
d.add_heading('Risposta sintetica  Kurzantwort',1)
d.add_paragraph('Sem.01 applica nel modello i principi centrali della SSV, ma non è un controllore certificato né direttamente installabile su strada. Tempi, matrice dei conflitti, hardware, diagnostica, privacy e autorizzazione richiedono una verifica specifica del luogo.')
d.add_paragraph('Sem.01 bildet zentrale SSV-Grundsätze ab, ist aber keine zertifizierte Strassensteuerung. Zeiten, Konfliktmatrix, Hardware, Diagnose, Datenschutz und Bewilligung sind standortspezifisch nachzuweisen.')
d.add_heading('Gerarchia dei riferimenti  Rang der Grundlagen',1)
d.add_paragraph('LEGGE = disposizione vincolante. TECNICA = norma VSS, direttiva ASTRA o calcolo progettuale. PROGETTO = barriera prudenziale Sem.01. Una scelta sicura può essere necessaria senza essere descritta letteralmente in un singolo articolo.')
d.add_heading('Matrice completa  Vollständige Zuordnung',1)
t=d.add_table(rows=1,cols=3); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False; borders(t); t.columns[0].width=Inches(.55); t.columns[1].width=Inches(2.85); t.columns[2].width=Inches(3.6)
for i,h in enumerate(['ID','Regola e implementazione','Base e valutazione']): t.cell(0,i).text=h; shade(t.cell(0,i),'17343B'); [setattr(r.font.color,'rgb',RGBColor(255,255,255)) for r in t.cell(0,i).paragraphs[0].runs]; [setattr(r,'bold',True) for r in t.cell(0,i).paragraphs[0].runs]
for rid,it,de,impl,ref,kind,note in R:
 c=t.add_row().cells; c[0].text=rid; c[1].text=f'{it}\n{de}\n\n{impl}'; c[2].text=f'{kind} — {ref}\n{note}'
 trPr=c[0]._tc.getparent().get_or_add_trPr(); trPr.append(OxmlElement('w:cantSplit'))
 if int(rid[1:])%2==0:
  for q in c: shade(q,'F1F5F6')
 for q in c: margin(q); q.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
 c[0].paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER
d.add_heading('Esito e lavoro residuo  Ergebnis und Restarbeiten',1)
for x in ['Acquisire le norme VSS vigenti e costruire una matrice dei requisiti verificabile.','Rilevare geometria, distanze, velocità, visibilità e flussi del sito.','Calcolare intertempi veicolari e pedonali; proteggere i parametri.','Definire hardware fail-safe, feedback lampade, watchdog, alimentazione, radio e sensori.','Eseguire analisi dei rischi, revisione indipendente, FAT, SAT e collaudo con l’autorità.','Per camera/radar: base giuridica, valutazione privacy ed elaborazione locale minimizzata.']: d.add_paragraph('•  '+x)
d.add_heading('Fonti  Quellen',1)
for a,b in S: p=d.add_paragraph('•  '); link(p,a,b)
d.add_heading('Limite della verifica  Abgrenzung',1); d.add_paragraph('Verifica tecnica preliminare, non parere legale, omologazione o autorizzazione. La conformità dipende dall’intero impianto, dal sito, dalle prove e dall’autorità competente.')
d.core_properties.title='Sem.01 Report 1.1 Matrice regole e norme svizzere'; d.core_properties.author='Progetto Sem.01'; d.save(OUT)
MD.parent.mkdir(exist_ok=True); L=['# Sem.01 Report 1.1 — Regole software e riferimenti svizzeri','', '> Verifica tecnica preliminare al 5 settembre 2026. Non è una certificazione né un’autorizzazione alla posa.','', '## Matrice','', '| ID | Regola IT / DE | Nel software | Base | Tipo e valutazione |','|---|---|---|---|---|']
clean=lambda x:x.replace('|','/').replace('\n',' ')
for rid,it,de,impl,ref,kind,note in R: L.append(f'| {rid} | {clean(it)} / {clean(de)} | {clean(impl)} | {clean(ref)} | **{kind}** — {clean(note)} |')
L+=['','## Conclusione','','Il simulatore implementa principi centrali della SSV ma non è un controllore certificato. Prima dell’hardware servono norme VSS vigenti, parametri del luogo, analisi dei rischi, prove indipendenti e autorizzazione.','','## Fonti','']+[f'- [{a}]({b})' for a,b in S]
MD.write_text('\n'.join(L)+'\n',encoding='utf-8'); print(OUT); print(MD)
